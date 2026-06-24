"""Generate real pdf / docx / png / jpeg KYC test fixtures from the plain-text samples.

    python tools/make_samples.py   # writes samples/generated/*

- passport.pdf       digital PDF (text layer)      -> exercises the pypdf path
- ssn_card.docx      Word document                 -> exercises the python-docx path
- ine_credencial.png rendered image                -> exercises the Tesseract OCR path
- utility_bill.jpg   rendered image                -> exercises the Tesseract OCR path
"""
from __future__ import annotations

import pathlib

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "samples"
OUT = ROOT / "samples" / "generated"

_FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/Helvetica.ttc",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "DejaVuSans.ttf",
]


def _font(size: int):
    from PIL import ImageFont

    for path in _FONT_CANDIDATES:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def make_pdf(text: str, path: pathlib.Path) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Courier", size=11)
    for line in text.splitlines():
        pdf.cell(0, 6, text=(line or " "), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(str(path))


def make_docx(text: str, path: pathlib.Path) -> None:
    import docx

    doc = docx.Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))


def make_image(text: str, path: pathlib.Path, fmt: str) -> None:
    from PIL import Image, ImageDraw

    lines = text.splitlines() or [" "]
    font = _font(30)
    pad, lh, width = 36, 46, 1100
    height = pad * 2 + lh * len(lines)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((pad, pad + i * lh), line, fill="black", font=font)
    if fmt == "JPEG":
        img.save(str(path), "JPEG", quality=95)
    else:
        img.save(str(path), "PNG")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_pdf((SRC / "passport_specimen.txt").read_text(), OUT / "passport.pdf")
    make_docx((SRC / "us_ssn_card.txt").read_text(), OUT / "ssn_card.docx")
    make_image((SRC / "mx_ine_credencial.txt").read_text(), OUT / "ine_credencial.png", "PNG")
    make_image((SRC / "us_utility_bill.txt").read_text(), OUT / "utility_bill.jpg", "JPEG")
    for p in sorted(OUT.iterdir()):
        print(f"  {p.relative_to(ROOT)}  ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
